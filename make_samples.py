"""Generate sample_data/: a sample survey, the corrected workbook and the audit report."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from core import TitleMatcher, build_audit_report, inspect_snapshot, load_snapshot, parse_pairs, process_workbook

OUT = Path(__file__).parent / "sample_data"
OUT.mkdir(exist_ok=True)
SHOWS = ["Bigg Boss", "Kapil Sharma Show", "Mirzapur", "Panchayat", "Squid Game", "Kaun Banega Crorepati"]
SRC = ["Q2", "Q6", "Q14"]; DST = ["Q3", "Q7", "Q15"]; ORDER = ["Q2", "Q3", "Q6", "Q7", "Q14", "Q15"]

# (Q2, Q6, Q14, destination row pattern for Q3/Q7/Q15)
RESP = [
    (["Bigg Boss", "Mirzapur"], ["The Kapil Sharma Show"], ["Panchayat"], [1, 1, 1, 1, 1, 1]),   # exact + normalized
    (["Big Boss", "Mirzapurr"], [], ["Squid Game"], [1, 2, 1, 1, 0, 1]),                          # alias, fuzzy, a 2 and a 0
    (["Bigg Boss 20"], ["KBC"], [], [1, 1, 1, 1, 1, None]),                                       # season-only -> review; KBC alias
    (["Paatal Lok"], ["Sacred Games"], ["Money Heist"], [1, 1, 1, 1, 1, 1]),                      # unrelated -> nothing
    (["Panchayat", "Squid Game"], ["Panchayat"], [], [1, 1, 1, 2, 1, "n/a"]),                     # duplicates + 2 + other
]
def build(labels, name, pairs_text):
    wb = Workbook(); ws = wb.active; ws.title = "Survey Data"
    col = 2; cols = {}
    for q in ORDER:
        n = 4 if q in SRC else len(SHOWS)
        cols[q] = list(range(col, col + n)); col += n
    ws.cell(row=1, column=1, value="Survey export (sample)")
    ws.cell(row=2, column=1, value="Resp ID")
    for q in ORDER:
        ws.cell(row=2, column=cols[q][0], value=f"{labels[q]}. Which OTT shows have you watched?")
        ws.merge_cells(start_row=2, start_column=cols[q][0], end_row=2, end_column=cols[q][-1])
        for j, c in enumerate(cols[q]):
            ws.cell(row=3, column=c, value=f"{labels[q]}_{j+1}" if q in SRC else SHOWS[j])
    for i, (q2, q6, q14, pat) in enumerate(RESP):
        r = 4 + i
        ws.cell(row=r, column=1, value=5000 + i)
        for q, titles in zip(SRC, (q2, q6, q14)):
            for j, t in enumerate(titles):
                ws.cell(row=r, column=cols[q][j], value=t)
        for q in DST:
            for j, v in enumerate(pat):
                if v is not None:
                    ws.cell(row=r, column=cols[q][j], value=v)
    ws.column_dimensions["A"].width = 12
    wb.create_sheet("Notes")["A1"] = "This sheet must stay untouched"
    wb.save(OUT / f"{name}.xlsx")

    checks = parse_pairs(pairs_text) if pairs_text else None
    data = (OUT / f"{name}.xlsx").read_bytes()
    layout = inspect_snapshot(load_snapshot(data, "Survey Data"), checks)
    assert layout.confident, layout.issues
    res = process_workbook(data, f"{name}.xlsx", layout, TitleMatcher(), checks=checks)
    assert res.verified
    (OUT / f"{name}_corrected.xlsx").write_bytes(res.corrected_bytes)
    (OUT / f"{name}_audit_report.xlsx").write_bytes(build_audit_report(res))
    print(name, "changed:", res.totals["changed"], "review:", res.totals["review_total"])
    for c in res.changed_cells: print(" ", c["cell"], c["color"], c["match_type"], repr(c["source_title"]), "->", c["header"])


build({q: q for q in ORDER}, "sample_survey", None)
# Same survey, numbered differently (pairs to type in the app: see sample_data/README_renumbered.txt)
RENUM = {"Q2": "Q4", "Q3": "Q9", "Q6": "Q5", "Q7": "Q12", "Q14": "Q20", "Q15": "Q21"}
PAIRS = "Q4 -> Q9\nQ4 -> Q12\nQ5 -> Q12\nQ4 -> Q21\nQ5 -> Q21\nQ20 -> Q21"
build(RENUM, "sample_survey_renumbered", PAIRS)
(OUT / "README_renumbered.txt").write_text(
    "sample_survey_renumbered.xlsx uses Q4, Q9, Q5, Q12, Q20, Q21.\n"
    "In the app sidebar, 'Question pairs to check', enter:\n\n" + PAIRS + "\n")
