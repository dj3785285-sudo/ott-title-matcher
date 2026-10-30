"""Small helpers for Excel column notation."""
from __future__ import annotations

import re
from typing import Iterable, List

from openpyxl.utils import column_index_from_string, get_column_letter

from .errors import ColumnSpecError

MAX_EXCEL_COLUMN = 16384  # XFD


def col_letter(index: int) -> str:
    return get_column_letter(index)


def addr(row: int, col: int) -> str:
    """1-based (row, col) -> 'BY19'."""
    return f"{get_column_letter(col)}{row}"


def col_to_index(token: str) -> int:
    t = token.strip().replace("$", "")
    if not re.fullmatch(r"[A-Za-z]{1,3}", t):
        raise ColumnSpecError(
            f"'{token}' is not a valid column. Use column letters only, e.g. BK or BK:BQ."
        )
    try:
        idx = column_index_from_string(t.upper())
    except ValueError as exc:  # pragma: no cover - defensive
        raise ColumnSpecError(f"'{token}' is not a valid Excel column.") from exc
    if idx > MAX_EXCEL_COLUMN:
        raise ColumnSpecError(f"Column '{token}' is beyond Excel's last column (XFD).")
    return idx


def parse_columns(spec: str | None) -> List[int]:
    """Parse 'BK:BQ', 'BK,BL,BM' or a mix ('BK:BM, BO') into 1-based column indexes.

    Order of first appearance is kept and duplicates are removed.
    """
    if spec is None:
        return []
    tokens = [t for t in re.split(r"[,\s;]+", str(spec).strip()) if t]
    out: List[int] = []
    seen = set()
    for tok in tokens:
        if ":" in tok:
            a, _, b = tok.partition(":")
            ia, ib = col_to_index(a), col_to_index(b)
            if ia > ib:
                ia, ib = ib, ia
            cols: Iterable[int] = range(ia, ib + 1)
        else:
            cols = [col_to_index(tok)]
        for c in cols:
            if c not in seen:
                seen.add(c)
                out.append(c)
    return out


def columns_to_spec(columns: Iterable[int]) -> str:
    """[63, 64, 65, 67] -> 'BK:BM,BO'."""
    cols = sorted(set(columns))
    if not cols:
        return ""
    parts = []
    start = prev = cols[0]
    for c in cols[1:]:
        if c == prev + 1:
            prev = c
            continue
        parts.append((start, prev))
        start = prev = c
    parts.append((start, prev))
    return ",".join(
        get_column_letter(a) if a == b else f"{get_column_letter(a)}:{get_column_letter(b)}"
        for a, b in parts
    )
