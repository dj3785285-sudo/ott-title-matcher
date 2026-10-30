"""Respondent-level matching and correction of destination cells.

Pipeline
--------
1. ``analyze``   - for each configured check (default: the six standard pairs), read every respondent's source titles
                   and resolve them against the destination headers.
2. ``decide``    - group matches by destination cell (across all checks), pick one
                   deterministic winner per cell (exact > normalized > alias > fuzzy,
                   then higher score, then lower check number) and apply the cell rules.
3. ``apply``     - write 0 + fill colour into a *copy* of the workbook (OpenPyXL).
4. ``verify``    - re-open the saved bytes and prove that only the intended cells differ.

The uploaded bytes are never modified; a new workbook is produced.
"""
from __future__ import annotations

import gc
import hashlib
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from openpyxl.styles import PatternFill

from .errors import MappingError
from .excel_utils import addr
from .pairs import DEFAULT_CHECKS, destinations_of, sources_of
from .inspector import (
    SheetLayout,
    dest_headers,
    validate_layout,
)
from .matching import (
    AUTO_TYPES,
    COLOR_FOR,
    EXACT,
    FUZZY,
    ALIAS,
    NORMALIZED,
    PRIORITY,
    Candidate,
    scorer_name,
    HeaderSet,
    TitleMatcher,
)
from .workbook_io import (
    Snapshot,
    detect_unsupported_features,
    is_blank,
    load_snapshot,
    open_workbook,
    snapshot_from_sheet,
)

# (check number, source question, destination question) - exactly the six required checks
CHECKS: List[Tuple[int, str, str]] = list(DEFAULT_CHECKS)  # default; callers may pass their own

YELLOW_HEX = "FFFF00"
GREEN_HEX = "00B050"
COLOR_HEX = {"Yellow": YELLOW_HEX, "Green": GREEN_HEX}

# Several titles in one source cell may be separated by these characters.
SPLIT_RE = re.compile(r"[\n\r;|]+")

PRIORITY_RULE_TEXT = (
    "When several source titles (or several checks) point at the same destination cell, the cell is "
    "changed and counted once. The winning match is chosen deterministically: exact beats normalized, "
    "normalized beats recognised-variant (alias), alias beats fuzzy; ties go to the higher similarity "
    "score, then the lower check number, then the leftmost source column. The cell colour follows the "
    "winning match (yellow = exact, green = anything else). Other matches to that cell are logged as "
    "'duplicate' and not counted as changes."
)


# --------------------------------------------------------------------------
# Value classification
# --------------------------------------------------------------------------
ONE, TWO, ZERO, BLANK, FORMULA, OTHER = "one", "two", "zero", "blank", "formula", "other"


def classify_value(v) -> str:
    """Classify a destination cell value according to the correction rules."""
    if v is None:
        return BLANK
    if isinstance(v, bool):
        return OTHER
    if isinstance(v, (int, float)):
        if v == 1:
            return ONE
        if v == 2:
            return TWO
        if v == 0:
            return ZERO
        return OTHER
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return BLANK
        if s.startswith("="):
            return FORMULA
        return {"1": ONE, "2": TWO, "0": ZERO}.get(s, OTHER)
    return OTHER


# --------------------------------------------------------------------------
# Data classes
# --------------------------------------------------------------------------
@dataclass
class MatchEvent:
    check_no: int
    src_q: str
    dst_q: str
    row: int
    src_col: int
    src_title: str
    dst_col: int
    dst_header: str
    type: str
    score: float
    reason: str
    seq: int = 0


@dataclass
class ReviewEvent:
    check_no: int
    src_q: str
    dst_q: str
    row: int
    src_col: int
    src_title: str
    status: str  # 'ambiguous' | 'review'
    candidates: List[Candidate]
    reason: str


@dataclass
class Analysis:
    events: List[MatchEvent] = field(default_factory=list)
    reviews: List[ReviewEvent] = field(default_factory=list)
    titles_checked: Counter = field(default_factory=Counter)
    header_cells: Dict[str, Dict[int, str]] = field(default_factory=dict)


@dataclass
class Decision:
    row: int
    col: int
    winner: MatchEvent
    others: List[MatchEvent]
    original: object
    kind: str
    action: str
    final: object
    color: str  # '' when not changed


@dataclass
class ProcessResult:
    filename: str
    sheet: str
    processed_at: datetime
    threshold: float
    layout: SheetLayout
    corrected_bytes: bytes
    audit_rows: List[dict]
    review_rows: List[dict]
    per_check: List[dict]
    totals: dict
    changed_cells: List[dict]
    feature_warnings: List[str]
    layout_warnings: List[str]
    verification_issues: List[str]
    verified: bool
    scorer: str = ""
    checks: List[Tuple[int, str, str]] = field(default_factory=lambda: list(CHECKS))


# --------------------------------------------------------------------------
# Step 1: analysis
# --------------------------------------------------------------------------
def split_titles(value: str) -> List[str]:
    return [p.strip() for p in SPLIT_RE.split(value) if p.strip()]


def analyze(snap: Snapshot, layout: SheetLayout, matcher: TitleMatcher, checks=None) -> Analysis:
    checks = list(checks if checks is not None else CHECKS)
    out = Analysis()
    header_sets: Dict[str, HeaderSet] = {}
    for qid in {dq for _, _, dq in checks}:
        items = dest_headers(snap, layout, qid)
        header_sets[qid] = HeaderSet(items)
        out.header_cells[qid] = {c: t for c, t in items}

    seq = 0
    for check_no, sq, dq in checks:
        hs = header_sets[dq]
        src_cols = layout.cols(sq)
        for row in range(layout.first_row, layout.last_row + 1):
            for sc in src_cols:
                value = snap.value(row, sc)
                if not isinstance(value, str) or not value.strip():
                    continue
                for title in split_titles(value):
                    out.titles_checked[check_no] += 1
                    res = matcher.resolve(title, hs)
                    if res.status == "matched":
                        cand = res.best
                        seq += 1
                        out.events.append(
                            MatchEvent(
                                check_no, sq, dq, row, sc, title, cand.col, cand.header,
                                cand.result.type, cand.result.score, cand.result.reason, seq,
                            )
                        )
                    elif res.status in ("ambiguous", "review"):
                        out.reviews.append(
                            ReviewEvent(
                                check_no, sq, dq, row, sc, title, res.status,
                                list(res.candidates), res.reason,
                            )
                        )
    return out


# --------------------------------------------------------------------------
# Step 2: decisions
# --------------------------------------------------------------------------
def _winner_key(ev: MatchEvent):
    return (-PRIORITY[ev.type], -ev.score, ev.check_no, ev.src_col, ev.seq)


def decide(snap: Snapshot, analysis: Analysis) -> List[Decision]:
    groups: Dict[Tuple[int, int], List[MatchEvent]] = defaultdict(list)
    for ev in analysis.events:
        groups[(ev.row, ev.dst_col)].append(ev)

    decisions: List[Decision] = []
    for (row, col), evs in sorted(groups.items()):
        evs.sort(key=_winner_key)
        winner, others = evs[0], evs[1:]
        original = snap.raw_value(row, col)
        kind = classify_value(original)
        color = ""
        final = original
        if kind == ONE:
            final = "0" if isinstance(original, str) else 0
            color = COLOR_FOR[winner.type]
            action = "Changed 1 to 0"
        elif kind == TWO:
            action = "Preserved (value 2)"
        elif kind == ZERO:
            action = "No change (already 0)"
        elif kind == BLANK:
            action = "Skipped (blank cell)"
        elif kind == FORMULA:
            action = "Skipped (cell contains a formula)"
        else:
            action = "Skipped (other value)"
        decisions.append(Decision(row, col, winner, others, original, kind, action, final, color))
    return decisions


# --------------------------------------------------------------------------
# Reporting rows / statistics
# --------------------------------------------------------------------------
def _display(v):
    return "" if v is None else v


def build_audit_rows(layout: SheetLayout, decisions: List[Decision]) -> List[dict]:
    rows: List[dict] = []
    for d in decisions:
        for ev in [d.winner, *d.others]:
            is_winner = ev is d.winner
            if is_winner:
                action = d.action
                modified = d.kind == ONE
                color = d.color
            else:
                action = f"Duplicate match - cell already handled by Check {d.winner.check_no}"
                modified = False
                color = ""
            rows.append(
                {
                    "Worksheet": layout.sheet,
                    "Check": ev.check_no,
                    "Source Question": ev.src_q,
                    "Destination Question": ev.dst_q,
                    "Respondent Row": ev.row,
                    "Source Title": ev.src_title,
                    "Source Cell": addr(ev.row, ev.src_col),
                    "Destination Title Header": ev.dst_header,
                    "Destination Header Cell": addr(layout.header_row, ev.dst_col),
                    "Destination Cell": addr(d.row, d.col),
                    "Original Value": _display(d.original),
                    "Final Value": _display(d.final),
                    "Match Type": ev.type,
                    "Match Confidence": round(ev.score, 2),
                    "Cell Modified": "Yes" if modified else "No",
                    "Action Taken": action,
                    "Highlight Color": color,
                }
            )
    rows.sort(key=lambda r: (r["Check"], r["Respondent Row"], r["Source Cell"], r["Destination Cell"]))
    return rows


def build_review_rows(layout: SheetLayout, analysis: Analysis) -> List[dict]:
    agg: Dict[tuple, dict] = {}
    for rv in analysis.reviews:
        status = "Ambiguous" if rv.status == "ambiguous" else "Near match"
        for cand in rv.candidates:
            key = (rv.check_no, rv.src_title.strip(), cand.col)
            row = agg.get(key)
            if row is None:
                row = agg[key] = {
                    "Check": rv.check_no,
                    "Source Question": rv.src_q,
                    "Destination Question": rv.dst_q,
                    "Source Title": rv.src_title,
                    "Source Cells": [],
                    "Occurrences": 0,
                    "Destination Header": cand.header,
                    "Destination Header Cell": addr(layout.header_row, cand.col),
                    "Similarity Score": round(cand.result.score, 2),
                    "Candidate Type": cand.result.type,
                    "Status": status,
                    "Reason for Manual Review": (
                        rv.reason if rv.status == "ambiguous" else cand.result.reason
                    ),
                }
            row["Occurrences"] += 1
            a = addr(rv.row, rv.src_col)
            if len(row["Source Cells"]) < 10 and a not in row["Source Cells"]:
                row["Source Cells"].append(a)
    rows = []
    for key in sorted(agg, key=lambda k: (k[0], k[1].casefold(), k[2])):
        r = agg[key]
        cells = r["Source Cells"]
        more = "" if r["Occurrences"] <= len(cells) else ", ..."
        r["Source Cells"] = ", ".join(cells) + more
        rows.append(r)
    return rows


def build_statistics(analysis: Analysis, decisions: List[Decision], checks=None) -> Tuple[List[dict], dict]:
    checks = list(checks if checks is not None else CHECKS)
    per: Dict[int, Counter] = {n: Counter() for n, _, _ in checks}
    for n, c in analysis.titles_checked.items():
        per[n]["titles_checked"] = c
    for ev in analysis.events:
        per[ev.check_no][ev.type] += 1
    for rv in analysis.reviews:
        per[rv.check_no]["ambiguous" if rv.status == "ambiguous" else "near_match"] += 1
    for d in decisions:
        c = per[d.winner.check_no]
        c["unique_cells_matched"] += 1
        c["duplicates_suppressed"] += len(d.others)
        c[{ONE: "changed", TWO: "preserved_2", ZERO: "already_0", BLANK: "blank",
           FORMULA: "other", OTHER: "other"}[d.kind]] += 1
        if d.kind == FORMULA:
            c["formula_skipped"] += 1

    keys = [
        "titles_checked", EXACT, NORMALIZED, ALIAS, FUZZY, "ambiguous", "near_match",
        "unique_cells_matched", "changed", "preserved_2", "already_0", "blank", "other",
        "duplicates_suppressed", "formula_skipped",
    ]
    rows = []
    totals = Counter()
    for n, sq, dq in checks:
        row = {"check": n, "source": sq, "destination": dq}
        for k in keys:
            row[k] = int(per[n][k])
            totals[k] += int(per[n][k])
        row["review_total"] = row["ambiguous"] + row["near_match"]
        rows.append(row)
    tot = {k: int(totals[k]) for k in keys}
    tot["review_total"] = tot["ambiguous"] + tot["near_match"]
    tot["title_matches"] = tot[EXACT] + tot[NORMALIZED] + tot[ALIAS] + tot[FUZZY]
    return rows, tot


# --------------------------------------------------------------------------
# Step 3 + 4: write and verify
# --------------------------------------------------------------------------
def _fill(color_name: str) -> PatternFill:
    hexv = COLOR_HEX[color_name]
    return PatternFill(fill_type="solid", start_color=hexv, end_color=hexv)


def _rgb_tail(cell) -> str:
    try:
        rgb = cell.fill.fgColor.rgb
        return rgb[-6:].upper() if isinstance(rgb, str) else ""
    except Exception:  # pragma: no cover
        return ""


def _cell_signature(c) -> tuple:
    """(value, fill type, fill rgb, number format) - works for normal and read-only cells."""
    fill = c.fill
    ftype = getattr(fill, "fill_type", None) if fill is not None else None
    return (c.value, ftype, _rgb_tail(c), c.number_format or "General")


def _fingerprints(data: bytes, skip: Dict[Tuple[str, int, int], str]):
    """Stream one workbook (read-only mode, low memory).

    Returns sheet names, per-sheet (hash of every cell except ``skip``, dimensions) and the
    (value, fill) of the skipped cells.
    """
    wb = open_workbook(data, read_only=True)
    names = list(wb.sheetnames)
    sheets: Dict[str, Tuple[str, Tuple[int, int]]] = {}
    special: Dict[Tuple[str, int, int], tuple] = {}
    for ws in wb.worksheets:
        h = hashlib.sha256()
        n_rows = n_cols = 0
        for r, row in enumerate(ws.iter_rows(), start=1):
            n_rows = r
            for c_idx, c in enumerate(row, start=1):
                n_cols = max(n_cols, c_idx)
                value, ftype, rgb, numfmt = _cell_signature(c)
                key = (ws.title, r, c_idx)
                if key in skip:
                    special[key] = (value, rgb)
                    continue
                if value is None and ftype in (None, "none") and numfmt == "General":
                    continue  # empty and unstyled: identical whether stored or not
                h.update(repr((r, c_idx, value, ftype, rgb, numfmt)).encode())
        sheets[ws.title] = (h.hexdigest(), (n_rows, n_cols))
    wb.close()
    return names, sheets, special


def _detailed_diff(original: bytes, output: bytes, changed) -> List[str]:
    """Slow path used only after a fingerprint mismatch: list the differing cells."""
    issues: List[str] = []
    wo, wn = open_workbook(original), open_workbook(output)
    for name in wo.sheetnames:
        so, sn = wo[name], wn[name]
        for row in so.iter_rows():
            for co in row:
                if (name, co.row, co.column) in changed:
                    continue
                cn = sn.cell(row=co.row, column=co.column)
                if co.value != cn.value:
                    issues.append(f"{name}!{co.coordinate}: value changed unexpectedly")
                elif (co.fill.fill_type, _rgb_tail(co)) != (cn.fill.fill_type, _rgb_tail(cn)):
                    issues.append(f"{name}!{co.coordinate}: fill changed unexpectedly")
                elif co.number_format != cn.number_format:
                    issues.append(f"{name}!{co.coordinate}: number format changed unexpectedly")
                if len(issues) >= 50:
                    return issues + ["... more discrepancies not listed"]
    return issues


def verify_output(original: bytes, output: bytes, changed: Dict[Tuple[str, int, int], str]) -> List[str]:
    """Reopen both workbooks and prove that only the intended cells differ.

    Returns a list of discrepancies (empty list = verified).
    """
    issues: List[str] = []
    names_o, sheets_o, special_o = _fingerprints(original, changed)
    names_n, sheets_n, special_n = _fingerprints(output, changed)
    if names_o != names_n:
        return [f"Worksheet names/order differ: {names_o} vs {names_n}"]
    for key, color in changed.items():
        label = f"{key[0]}!{addr(key[1], key[2])}"
        if classify_value(special_o.get(key, (None,))[0]) != ONE:
            issues.append(f"{label}: changed although the original value was not 1")
        new_val, new_fill = special_n.get(key, (None, ""))
        if classify_value(new_val) != ZERO:
            issues.append(f"{label}: new value is not 0")
        if new_fill != COLOR_HEX[color]:
            issues.append(f"{label}: highlight colour missing")
    for name in names_o:
        if sheets_o[name][0] != sheets_n[name][0]:
            issues.extend(_detailed_diff(original, output, changed))
            break
    return issues[:51]


def process_workbook(
    data: bytes,
    filename: str,
    layout: SheetLayout,
    matcher: TitleMatcher,
    *,
    verify: bool = True,
    processed_at: Optional[datetime] = None,
    checks=None,
) -> ProcessResult:
    checks = list(checks if checks is not None else CHECKS)
    wb = open_workbook(data)  # the only full load: used for reading AND for writing the copy
    if layout.sheet not in wb.sheetnames:
        raise MappingError(f"Worksheet '{layout.sheet}' was not found in the workbook.")
    ws = wb[layout.sheet]
    snap = snapshot_from_sheet(ws, data)
    errors, layout_warnings = validate_layout(snap, layout, checks)
    if errors:
        raise MappingError(errors)

    analysis = analyze(snap, layout, matcher, checks)
    decisions = decide(snap, analysis)

    # --- safety net: never touch a source-question column --------------------------
    src_cols = {c for q in sources_of(checks) for c in layout.cols(q)}
    for d in decisions:
        if d.col in src_cols:  # pragma: no cover - validate_layout already forbids overlaps
            raise MappingError(f"Internal safety stop: {addr(d.row, d.col)} is a source-question cell.")

    changed: Dict[Tuple[str, int, int], str] = {}
    changed_cells: List[dict] = []
    for d in decisions:
        if d.kind != ONE:
            continue
        cell = ws.cell(row=d.row, column=d.col)
        cell.value = d.final
        cell.fill = _fill(d.color)
        changed[(layout.sheet, d.row, d.col)] = d.color
        changed_cells.append(
            {
                "cell": addr(d.row, d.col),
                "row": d.row,
                "col": d.col,
                "color": d.color,
                "match_type": d.winner.type,
                "source_title": d.winner.src_title,
                "header": d.winner.dst_header,
            }
        )
    buf = io.BytesIO()
    wb.save(buf)
    corrected = buf.getvalue()
    del wb, ws, snap  # free memory before verification re-opens the files
    gc.collect()

    issues: List[str] = []
    if verify:
        issues = verify_output(data, corrected, changed)

    audit_rows = build_audit_rows(layout, decisions)
    review_rows = build_review_rows(layout, analysis)
    per_check, totals = build_statistics(analysis, decisions, checks)

    return ProcessResult(
        filename=filename,
        sheet=layout.sheet,
        processed_at=processed_at or datetime.now(),
        threshold=matcher.threshold,
        layout=layout,
        corrected_bytes=corrected,
        audit_rows=audit_rows,
        review_rows=review_rows,
        per_check=per_check,
        totals=totals,
        changed_cells=changed_cells,
        feature_warnings=detect_unsupported_features(data),
        layout_warnings=layout_warnings,
        verification_issues=issues,
        verified=verify and not issues,
        scorer=scorer_name(),
        checks=checks,
    )


# --------------------------------------------------------------------------
# Preview (dry run - nothing is written)
# --------------------------------------------------------------------------
def preview_matches(data: bytes, layout: SheetLayout, matcher: TitleMatcher, checks=None) -> dict:
    """Dry run: distinct title matches, review items and how many cells would change."""
    checks = list(checks if checks is not None else CHECKS)
    snap = load_snapshot(data, layout.sheet)
    errors, warnings = validate_layout(snap, layout, checks)
    if errors:
        raise MappingError(errors)
    analysis = analyze(snap, layout, matcher, checks)
    decisions = decide(snap, analysis)
    _, totals = build_statistics(analysis, decisions, checks)

    agg: Dict[tuple, dict] = {}
    for ev in analysis.events:
        key = (ev.check_no, ev.src_title.strip().casefold(), ev.dst_col, ev.type)
        r = agg.get(key)
        if r is None:
            r = agg[key] = {
                "Check": f"{ev.check_no}: {ev.src_q} -> {ev.dst_q}",
                "Source Title": ev.src_title,
                "Destination Header": ev.dst_header,
                "Classification": ev.type,
                "Colour if changed": COLOR_FOR[ev.type],
                "Score": round(ev.score, 1),
                "Occurrences": 0,
            }
        r["Occurrences"] += 1
    matches = sorted(agg.values(), key=lambda r: (r["Check"], -PRIORITY[r["Classification"]], r["Source Title"].casefold()))
    review = build_review_rows(layout, analysis)
    return {
        "matches": matches,
        "review": review,
        "totals": totals,
        "would_change": totals["changed"],
        "warnings": warnings,
    }
