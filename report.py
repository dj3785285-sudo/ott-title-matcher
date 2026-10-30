"""Excel audit report (Summary / Audit Log / Review Matches)."""
from __future__ import annotations

import io
from typing import List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .excel_utils import columns_to_spec
from .pairs import destinations_of, questions_of, sources_of
from .matching import scorer_name
from .processor import (
    COLOR_HEX,
    PRIORITY_RULE_TEXT,
    ProcessResult,
)

HEADER_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AUDIT_COLUMNS = [
    "Worksheet", "Check", "Source Question", "Destination Question", "Respondent Row",
    "Source Title", "Source Cell", "Destination Title Header", "Destination Header Cell",
    "Destination Cell", "Original Value", "Final Value", "Match Type", "Match Confidence",
    "Cell Modified", "Action Taken", "Highlight Color",
]
REVIEW_COLUMNS = [
    "Check", "Source Question", "Destination Question", "Source Title", "Source Cells",
    "Occurrences", "Destination Header", "Destination Header Cell", "Similarity Score",
    "Candidate Type", "Status", "Reason for Manual Review",
]

PER_CHECK_COLUMNS = [
    ("check", "Check"), ("source", "Source Q"), ("destination", "Destination Q"),
    ("titles_checked", "Source titles checked"),
    ("exact", "Exact matches"), ("normalized", "Normalized matches"),
    ("alias", "Recognised-variant (alias) matches"), ("fuzzy", "Fuzzy matches"),
    ("changed", "Cells changed 1 to 0"), ("preserved_2", "Cells preserved (value 2)"),
    ("already_0", "Cells already 0"), ("blank", "Blank cells skipped"),
    ("other", "Other values skipped"), ("duplicates_suppressed", "Duplicate matches suppressed"),
    ("ambiguous", "Ambiguous (manual review)"), ("near_match", "Near matches (manual review)"),
]


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def _write_summary(ws, res: ProcessResult) -> None:
    ws.sheet_view.showGridLines = False
    r = 1
    ws.cell(row=r, column=1, value="OTT Title Matching - Audit Report").font = Font(bold=True, size=16, color="1F3864")
    r += 2

    info = [
        ("Original filename", res.filename),
        ("Processing date and time", res.processed_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Selected worksheet", res.sheet),
        ("Fuzzy-matching threshold", f"{res.threshold:g}%"),
        ("Similarity scorer", res.scorer or scorer_name()),
        ("Question-label row", res.layout.label_row),
        ("Title-header row", res.layout.header_row),
        ("First respondent row", res.layout.first_row),
        ("Last respondent row", res.layout.last_row),
        ("Output verification", "PASSED - re-opened the saved workbook; only intended cells differ"
         if res.verified else "NOT VERIFIED / issues found: " + "; ".join(res.verification_issues[:5])),
    ]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Colour legend").font = SECTION_FONT
    r += 1
    legend = [
        ("Yellow", "Exact match, changed from 1 to 0."),
        ("Green", "Normalized or fuzzy title match (incl. recognised variants), changed from 1 to 0."),
    ]
    for name, text in legend:
        c = ws.cell(row=r, column=1, value=name)
        c.fill = PatternFill("solid", start_color=COLOR_HEX[name], end_color=COLOR_HEX[name])
        c.font = Font(bold=True, color="000000" if name == "Yellow" else "FFFFFF")
        c.border = BORDER
        ws.cell(row=r, column=2, value=f"{name} = {text}")
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Question-pair checks used").font = SECTION_FONT
    r += 1
    for n, s_, d_ in res.checks:
        ws.cell(row=r, column=1, value=f"Check {n}")
        ws.cell(row=r, column=2, value=f"{s_} (source) -> {d_} (destination)")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Question mappings used").font = SECTION_FONT
    r += 1
    for j, h in enumerate(["Question", "Role", "Columns", "Label cell"], start=1):
        ws.cell(row=r, column=j, value=h)
    _style_header(ws, r, 4)
    r += 1
    sources = set(sources_of(res.checks))
    for qid in questions_of(res.checks):
        q = res.layout.questions.get(qid)
        role = "Source (never modified)" if qid in sources else "Destination (1 -> 0 corrections)"
        vals = [qid, role, columns_to_spec(q.columns) if q else "", (q.label_cell if q else "") or "(manual)"]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v).border = BORDER
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Overall totals").font = SECTION_FONT
    r += 1
    t = res.totals
    overall = [
        ("Source titles checked", t["titles_checked"]),
        ("Exact matches", t["exact"]),
        ("Normalized matches", t["normalized"]),
        ("Recognised-variant (alias) matches", t["alias"]),
        ("Fuzzy matches", t["fuzzy"]),
        ("Cells changed from 1 to 0", t["changed"]),
        ("Cells preserved because value was 2", t["preserved_2"]),
        ("Cells already containing 0", t["already_0"]),
        ("Blank cells skipped", t["blank"]),
        ("Other values skipped (incl. formulas)", t["other"]),
        ("Duplicate matches suppressed (cell already handled)", t["duplicates_suppressed"]),
        ("Ambiguous matches requiring review", t["ambiguous"]),
        ("Near matches requiring review", t["near_match"]),
    ]
    for k, v in overall:
        ws.cell(row=r, column=1, value=k).border = BORDER
        c = ws.cell(row=r, column=2, value=v)
        c.border = BORDER
        c.alignment = Alignment(horizontal="left")
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Statistics per question-pair check").font = SECTION_FONT
    r += 1
    for j, (_, title) in enumerate(PER_CHECK_COLUMNS, start=1):
        ws.cell(row=r, column=j, value=title)
    _style_header(ws, r, len(PER_CHECK_COLUMNS))
    ws.row_dimensions[r].height = 45
    r += 1
    for row in res.per_check:
        for j, (key, _) in enumerate(PER_CHECK_COLUMNS, start=1):
            ws.cell(row=r, column=j, value=row[key]).border = BORDER
        r += 1
    for j, (key, _) in enumerate(PER_CHECK_COLUMNS, start=1):
        if j == 1:
            v = "Total"
        elif key in ("source", "destination"):
            v = ""
        else:
            v = res.totals[key]
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(bold=True)
        c.border = BORDER
    r += 2

    ws.cell(row=r, column=1, value="Rules applied").font = SECTION_FONT
    r += 1
    notes = [
        "Destination cell = 1 -> changed to 0 and highlighted. = 2 -> preserved. = 0, blank or any other value -> left unchanged.",
        "Source-question answers and unrelated cells are never modified.",
        "Ambiguous or low-confidence matches are listed in 'Review Matches' and are NOT applied.",
        PRIORITY_RULE_TEXT,
        "Match counts above count title matches (before de-duplication); cell counts are per unique destination cell.",
    ]
    for n in notes:
        c = ws.cell(row=r, column=1, value=n)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 15 * (1 + len(n) // 120)
        r += 1

    if res.feature_warnings or res.layout_warnings:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = SECTION_FONT
        r += 1
        for w in [*res.feature_warnings, *res.layout_warnings]:
            ws.cell(row=r, column=1, value=w)
            r += 1

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34
    for j in range(3, len(PER_CHECK_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16


def _format_table(ws, ncols: int, widths: List[int]) -> None:
    _style_header(ws, 1, ncols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32


def build_audit_report(res: ProcessResult) -> bytes:
    """Return the audit report workbook as bytes."""
    audit_df = pd.DataFrame(res.audit_rows, columns=AUDIT_COLUMNS)
    review_df = pd.DataFrame(res.review_rows, columns=REVIEW_COLUMNS)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        audit_df.to_excel(writer, sheet_name="Audit Log", index=False)
        review_df.to_excel(writer, sheet_name="Review Matches", index=False)
        wb = writer.book
        summary = wb.create_sheet("Summary", 0)
        _write_summary(summary, res)

        ws = wb["Audit Log"]
        _format_table(ws, len(AUDIT_COLUMNS), [14, 7, 10, 12, 11, 24, 11, 24, 12, 12, 10, 10, 12, 11, 10, 46, 12])
        color_col = AUDIT_COLUMNS.index("Highlight Color") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=color_col)
            if cell.value in COLOR_HEX:
                cell.fill = PatternFill("solid", start_color=COLOR_HEX[cell.value], end_color=COLOR_HEX[cell.value])

        ws = wb["Review Matches"]
        _format_table(ws, len(REVIEW_COLUMNS), [7, 10, 12, 26, 30, 11, 26, 14, 11, 12, 12, 60])
        if review_df.empty:
            ws.cell(row=2, column=1, value="No ambiguous or near-matching titles needed review.")
        wb.active = 0
    return buf.getvalue()
