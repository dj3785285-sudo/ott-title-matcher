"""Reading workbooks into lightweight snapshots + unsupported-feature warnings."""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .errors import WorkbookError

Merged = Tuple[int, int, int, int]  # min_row, min_col, max_row, max_col


@dataclass
class Snapshot:
    """In-memory copy of one worksheet (1-based accessors).

    ``values``   - what a person sees (cached formula results are used for titles).
    ``raw``      - what is stored in the file (formulas appear as '=...' strings);
                   destination cells are classified from this so formulas are never overwritten.
    """

    sheet: str
    values: List[list]
    raw: List[list]
    max_row: int
    max_col: int
    merged: List[Merged] = field(default_factory=list)

    def value(self, row: int, col: int):
        if 1 <= row <= self.max_row and 1 <= col <= self.max_col:
            return self.values[row - 1][col - 1]
        return None

    def raw_value(self, row: int, col: int):
        if 1 <= row <= self.max_row and 1 <= col <= self.max_col:
            return self.raw[row - 1][col - 1]
        return None

    def merged_at(self, row: int, col: int) -> Optional[Merged]:
        for m in self.merged:
            if m[0] <= row <= m[2] and m[1] <= col <= m[3]:
                return m
        return None


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _open(data: bytes, data_only: bool, read_only: bool = False):
    if not data or data[:2] != b"PK":
        raise WorkbookError(
            "This does not look like an .xlsx workbook. Please upload an Excel .xlsx file "
            "(older .xls files must be re-saved as .xlsx first)."
        )
    try:
        return load_workbook(io.BytesIO(data), data_only=data_only, read_only=read_only)
    except zipfile.BadZipFile as exc:
        raise WorkbookError("The file is corrupted or is not a valid .xlsx workbook.") from exc
    except Exception as exc:  # openpyxl raises many different errors
        raise WorkbookError(f"The workbook could not be read: {exc}") from exc


def open_workbook(data: bytes, data_only: bool = False, read_only: bool = False):
    return _open(data, data_only, read_only)


def sheet_names(data: bytes) -> List[str]:
    return list(_open(data, True).sheetnames)


def _grid(ws) -> List[list]:
    return [
        list(r)
        for r in ws.iter_rows(
            min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=True
        )
    ]


def _has_formula(grid: List[list]) -> bool:
    return any(isinstance(v, str) and v.startswith("=") for row in grid for v in row)


def snapshot_from_sheet(ws_e, data: bytes) -> Snapshot:
    """Build a snapshot from an already-loaded (formula) worksheet.

    ``raw`` holds what is stored; ``values`` equals ``raw`` except that formula cells show
    their cached results.  The cached-values workbook is only opened when the sheet
    really contains formulas, so ordinary survey exports are loaded once.
    """
    raw = _grid(ws_e)
    values = [row[:] for row in raw]
    if _has_formula(raw):
        try:
            wb_v = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            cached = [list(r) for r in wb_v[ws_e.title].iter_rows(values_only=True)]
            wb_v.close()
        except Exception:  # cached values are a convenience only
            cached = []
        for r, row in enumerate(values):
            for c, v in enumerate(row):
                if isinstance(v, str) and v.startswith("="):
                    got = cached[r][c] if r < len(cached) and c < len(cached[r]) else None
                    row[c] = got
    merged = [
        (m.min_row, m.min_col, m.max_row, m.max_col) for m in ws_e.merged_cells.ranges
    ]
    return Snapshot(ws_e.title, values, raw, ws_e.max_row, ws_e.max_column, merged)


def load_snapshot(data: bytes, sheet: str) -> Snapshot:
    wb = _open(data, False)
    if sheet not in wb.sheetnames:
        raise WorkbookError(f"Worksheet '{sheet}' was not found in the workbook.")
    return snapshot_from_sheet(wb[sheet], data)


def load_all_snapshots(data: bytes) -> Dict[str, Snapshot]:
    """Snapshots of every sheet (one workbook load), for inspection."""
    wb = _open(data, False)
    return {ws.title: snapshot_from_sheet(ws, data) for ws in wb.worksheets}


# --------------------------------------------------------------------------
# Unsupported-feature detection
# --------------------------------------------------------------------------
_FEATURES = [
    ("xl/charts/", "Charts are present. OpenPyXL cannot read charts, so they may be LOST in the corrected workbook."),
    ("xl/drawings/", "Drawings/shapes/images are present and may be lost or altered in the corrected workbook."),
    ("xl/media/", "Embedded images are present; they may be lost unless the Pillow library is available."),
    ("xl/vbaProject.bin", "A VBA macro project is present (macros are not kept in .xlsx output)."),
    ("xl/slicers/", "Slicers are present and may be lost."),
    ("xl/timelines/", "Timelines are present and may be lost."),
    ("xl/ctrlProps/", "Form controls are present and may be lost."),
    ("xl/embeddings/", "Embedded objects are present and may be lost."),
    ("xl/threadedComments/", "Threaded comments are present; they may be converted to plain notes or lost."),
    ("xl/pivotTables/", "Pivot tables are present; they are kept but should be checked after processing."),
]


def detect_unsupported_features(data: bytes) -> List[str]:
    """Return warnings about workbook parts that OpenPyXL may not round-trip."""
    warnings: List[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            for prefix, message in _FEATURES:
                if any(n.startswith(prefix) or n == prefix for n in names):
                    warnings.append(message)
            ext_hits = 0
            sparkline = False
            for n in names:
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n):
                    xml = zf.read(n)
                    if b"<extLst" in xml:
                        ext_hits += 1
                    if b"sparkline" in xml:
                        sparkline = True
            if sparkline:
                warnings.append("Sparklines are present and will be lost.")
            if ext_hits:
                warnings.append(
                    "Some worksheets use Excel extension features (e.g. newer conditional formatting "
                    "or data-validation options) that may be dropped."
                )
    except zipfile.BadZipFile:
        pass
    return warnings
