"""Workbook inspection, layouts and manual mapping (mandatory test 12)."""
import pytest

from core import (
    MappingError, build_layout, columns_to_spec, inspect_snapshot, load_snapshot,
    parse_columns, validate_layout,
)
from core.errors import ColumnSpecError
from core.inspector import ALL_QUESTIONS, question_number
from tests.helpers import SurveyBuilder, layout_for


def make(**kw):
    b = SurveyBuilder(**kw)
    for i in range(4):
        b.add(Q2=["Bigg Boss"], Q3=[1, 0, 2, None, 1][: len(b.headers)])
    return b


# ---- column notation ---------------------------------------------------------
def test_parse_columns_range_list_and_mix():
    assert parse_columns("BK:BM") == [63, 64, 65]
    assert parse_columns("BK,BL,BM") == [63, 64, 65]
    assert parse_columns("bk:bl, BO ; $BQ") == [63, 64, 67, 69]
    assert parse_columns("BM:BK") == [63, 64, 65]  # reversed range
    assert parse_columns("") == [] and parse_columns(None) == []


@pytest.mark.parametrize("bad", ["BK1:BQ9", "B-K", "12", "ABCD"])
def test_parse_columns_rejects_garbage(bad):
    with pytest.raises(ColumnSpecError):
        parse_columns(bad)


def test_columns_to_spec_roundtrip():
    assert columns_to_spec([63, 64, 65, 67]) == "BK:BM,BO"
    assert parse_columns(columns_to_spec([2, 3, 4, 9, 11, 12])) == [2, 3, 4, 9, 11, 12]


def test_question_number_variants():
    assert question_number("Q2") == 2
    assert question_number("Q2_1") == 2
    assert question_number("q 14. Which") == 14
    assert question_number("Question 15 - shows") == 15
    assert question_number("Quiz 2") is None
    assert question_number("Q") is None
    assert question_number(5) is None


# ---- detection in different layouts ----------------------------------------------
@pytest.mark.parametrize(
    "kw",
    [
        dict(),
        dict(top_pad=3),
        dict(left_pad=0),
        dict(left_pad=5, top_pad=1),
        dict(merged_labels=True),
        dict(gap=2, top_pad=2),
        dict(src_cols=5, merged_labels=True, top_pad=4),
        dict(headers=["A Show", "B Show", "C Show"], src_cols=2),
    ],
    ids=lambda kw: str(kw) or "default",
)
def test_detects_layout_variants(kw):
    b = make(**kw)
    layout = layout_for(b.build())
    assert layout.confident, layout.issues
    assert layout.label_row == b.label_row
    assert layout.header_row == b.header_row
    assert layout.first_row == b.first_row
    assert layout.last_row == b.row_of(3)
    for q in ALL_QUESTIONS:
        assert layout.cols(q) == b.cols[q], q


def test_merged_label_cells_detected():
    layout = layout_for(make(merged_labels=True).build())
    assert layout.questions["Q2"].merged and layout.questions["Q15"].merged


def test_different_sheet_name_and_extra_sheet_ignored():
    b = make(sheet="Raw Data 2025")
    layout = layout_for(b.build(), "Raw Data 2025")
    assert layout.sheet == "Raw Data 2025" and layout.confident


def test_missing_question_is_reported_not_guessed():
    b = make()
    data = b.build()
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Survey"]
    ws.cell(row=1, column=b.cols["Q7"][0]).value = "Something else"
    buf = io.BytesIO()
    wb.save(buf)
    layout = layout_for(buf.getvalue())
    assert not layout.confident
    assert any("Q7" in i for i in layout.issues)
    snap = load_snapshot(buf.getvalue(), "Survey")
    errors, _ = validate_layout(snap, layout)
    assert any("Q7" in e for e in errors)


def test_sheet_without_labels_needs_manual_mapping():
    from openpyxl import Workbook
    import io
    wb = Workbook()
    wb.active["A1"] = "hello"
    buf = io.BytesIO()
    wb.save(buf)
    layout = layout_for(buf.getvalue(), wb.active.title)
    assert not layout.confident and layout.found == 0


# ---- manual mapping / validation ---------------------------------------------------
def test_manual_layout_from_excel_notation():
    b = make(top_pad=2)
    data = b.build()
    specs = {q: columns_to_spec(b.cols[q]) for q in ALL_QUESTIONS}
    layout = build_layout("Survey", b.label_row, b.header_row, b.first_row, b.row_of(3), specs)
    errors, _ = validate_layout(load_snapshot(data, "Survey"), layout)
    assert errors == []


def test_overlapping_ranges_are_blocked():
    b = make()
    data = b.build()
    specs = {q: columns_to_spec(b.cols[q]) for q in ALL_QUESTIONS}
    specs["Q3"] = columns_to_spec(b.cols["Q2"] + b.cols["Q3"])  # Q3 swallows source Q2
    layout = build_layout("Survey", b.label_row, b.header_row, b.first_row, b.row_of(3), specs)
    errors, _ = validate_layout(load_snapshot(data, "Survey"), layout)
    assert any("mapped to both Q2 and Q3" in e for e in errors)


def test_bad_rows_are_blocked():
    b = make()
    data = b.build()
    specs = {q: columns_to_spec(b.cols[q]) for q in ALL_QUESTIONS}
    layout = build_layout("Survey", 1, 5, 3, 2, specs)
    errors, _ = validate_layout(load_snapshot(data, "Survey"), layout)
    assert any("below the title-header row" in e for e in errors)
    assert any("last respondent row" in e for e in errors)


def test_bad_column_spec_raises_mapping_error():
    with pytest.raises(MappingError):
        build_layout("Survey", 1, 2, 3, 4, {"Q2": "B1:C2"})


def test_text_subheader_row_is_skipped_with_issue():
    from openpyxl import load_workbook
    import io
    b = make()
    wb = load_workbook(io.BytesIO(b.build()))
    ws = wb["Survey"]
    ws.insert_rows(3)  # inserts a blank row; fill it with descriptive text
    for c in b.cols["Q3"]:
        ws.cell(row=3, column=c, value="Watched in last month")
    buf = io.BytesIO()
    wb.save(buf)
    layout = layout_for(buf.getvalue())
    assert layout.first_row == 4
    assert not layout.confident and any("Skipped" in i for i in layout.issues)
